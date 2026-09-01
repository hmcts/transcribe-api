"""Procedural Word frontmatter generator using python-docx and Pydantic."""

from __future__ import annotations

from pathlib import Path
from typing import TYPE_CHECKING

from docx import Document

from transcribe_api.documents.lib.typography import (
    add_blank_paragraphs,
    centered_crest,
    centered_text,
    label_value_line,
    left_block,
    normalize_paragraph_spacing,
    right_aligned_text,
    set_default_font,
)

if TYPE_CHECKING:
    from docx.document import Document as DocumentObject

    from transcribe_api.documents.lib.models import Frontmatter


def render_frontmatter(
    document: DocumentObject,
    frontmatter: Frontmatter,
    crest_path: str | Path | None = None,
) -> None:
    """Render frontmatter into an existing python-docx Document."""
    if crest_path:
        centered_crest(document, Path(crest_path))

    left_block(
        document,
        text=f"{frontmatter.tribunal_title}\n{frontmatter.tribunal_chamber}",
    )
    right_aligned_text(
        document,
        text=f"Appeal Number: {frontmatter.appeal_number}",
    )
    add_blank_paragraphs(document)
    centered_text(document, frontmatter.acts_heading, underline=True)
    add_blank_paragraphs(document)
    left_block(
        document,
        text=f"{frontmatter.hearing_location_prefix} {frontmatter.hearing_location}\n"
        f"{frontmatter.hearing_date_prefix} {frontmatter.hearing_date}",
    )
    add_blank_paragraphs(document)

    centered_text(document, frontmatter.before_heading)
    add_blank_paragraphs(document)
    centered_text(
        document,
        f"{frontmatter.judge_title} {frontmatter.judge_name}",
        all_caps=True,
    )
    add_blank_paragraphs(document)
    centered_text(document, frontmatter.between_heading)
    add_blank_paragraphs(document)
    centered_text(document, frontmatter.appellant_display_name())

    right_aligned_text(
        document,
        text=frontmatter.appellant_role_label,
        bold=False,
        underline=True,
    )
    centered_text(document, frontmatter.and_heading)
    add_blank_paragraphs(document)
    centered_text(document, frontmatter.respondent_name)
    right_aligned_text(
        document,
        text=frontmatter.respondent_role_label,
        bold=False,
        underline=True,
    )
    add_blank_paragraphs(document)

    representation = document.add_paragraph()
    normalize_paragraph_spacing(representation)
    heading_run = representation.add_run(frontmatter.representation_heading)
    heading_run.bold = True
    heading_run.underline = True
    add_blank_paragraphs(document)

    label_value_line(
        document,
        label=frontmatter.for_appellant_label,
        value=f"{frontmatter.appellant_placeholder} {frontmatter.appellant_representative}".strip(),
    )
    label_value_line(
        document,
        label=frontmatter.for_respondent_label,
        value=frontmatter.respondent_representative,
    )


def build_frontmatter(
    frontmatter: Frontmatter,
    output_path: str | Path,
    crest_path: str | Path | None = None,
) -> Path:
    """Build and save a frontmatter Word document procedurally."""
    document = Document()
    set_default_font(document)
    render_frontmatter(document, frontmatter, crest_path)

    destination = Path(output_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(destination))
    return destination


def build_frontmatter_document(
    frontmatter: Frontmatter,
    output_path: str | Path,
    crest_path: str | Path | None = None,
) -> Path:
    """Backward-compatible alias for build_frontmatter."""
    return build_frontmatter(frontmatter, output_path, crest_path)
