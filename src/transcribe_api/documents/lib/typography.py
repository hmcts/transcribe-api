from pathlib import Path

from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.shared import Inches, Pt

from transcribe_api.documents.lib.docx_style_profile import (
    DEFAULT_FONT_FAMILY,
    DEFAULT_FONT_SIZE_PT,
    apply_global_document_formatting,
)


def normalize_paragraph_spacing(paragraph) -> None:
    paragraph_format = paragraph.paragraph_format
    paragraph_format.space_before = Pt(0)
    paragraph_format.space_after = Pt(0)
    paragraph_format.line_spacing = 1.0
    paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE


def set_default_font(document: DocumentObject) -> None:
    style = document.styles["Normal"]
    style.font.name = DEFAULT_FONT_FAMILY
    style.font.size = Pt(DEFAULT_FONT_SIZE_PT)
    apply_global_document_formatting(document)


def centered_crest(document: DocumentObject, crest_path: Path) -> None:
    paragraph = document.add_paragraph()
    normalize_paragraph_spacing(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(crest_path), width=Inches(1.6))


def _aligned_text(
    document: DocumentObject,
    text: str,
    *,
    alignment: WD_ALIGN_PARAGRAPH,
    bold: bool = True,
    underline: bool = False,
    all_caps: bool = False,
) -> None:
    paragraph = document.add_paragraph()
    normalize_paragraph_spacing(paragraph)
    paragraph.alignment = alignment
    run = paragraph.add_run(text)
    run.bold = bold
    run.underline = underline
    run.font.all_caps = all_caps


def centered_text(
    document: DocumentObject,
    text: str,
    *,
    bold: bool = True,
    underline: bool = False,
    all_caps: bool = False,
) -> None:
    _aligned_text(document, text, alignment=WD_ALIGN_PARAGRAPH.CENTER, bold=bold, underline=underline, all_caps=all_caps)


def right_aligned_text(
    document: DocumentObject,
    text: str,
    *,
    bold: bool = True,
    underline: bool = False,
    all_caps: bool = False,
) -> None:
    _aligned_text(document, text, alignment=WD_ALIGN_PARAGRAPH.RIGHT, bold=bold, underline=underline, all_caps=all_caps)


def left_block(document: DocumentObject, text: str, *, bold: bool = True) -> None:
    paragraph = document.add_paragraph(text)
    normalize_paragraph_spacing(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if paragraph.runs:
        paragraph.runs[0].bold = bold


def label_value_line(
    document: DocumentObject,
    *,
    label: str,
    value: str,
    tab_position_pt: float = 130.0,
) -> None:
    paragraph = document.add_paragraph()
    normalize_paragraph_spacing(paragraph)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(tab_position_pt), WD_TAB_ALIGNMENT.LEFT)
    paragraph.add_run(label)
    paragraph.add_run(f"\t{value}")


def add_blank_paragraphs(document: DocumentObject, count: int = 1) -> None:
    for _ in range(max(count, 0)):
        paragraph = document.add_paragraph("")
        normalize_paragraph_spacing(paragraph)
