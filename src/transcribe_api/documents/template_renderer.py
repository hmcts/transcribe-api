"""Template rendering for hearing documents.

This module provides the interface for rendering Word documents from
templates using hearing form data and transcription content.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import TYPE_CHECKING

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from transcribe_api.documents.anonymity import (
    DELETE_MARKER,
    get_anonymity_content,
    parse_anonymity_status,
)
from transcribe_api.documents.content_loader import (
    get_hearing_description,
    get_hearing_title,
    get_legal_framework_content,
    get_legal_issues_display,
)
from transcribe_api.documents.hearing_document_router import build_hearing_document_from_document_data
from transcribe_api.documents.models import DocumentData, HearingFormData, TranscriptMessage
from transcribe_api.documents.styles import (
    INDENT_INCREMENT_CM,
    QUOTATION_BASE_INDENT_CM,
    STYLE_HEADING_1,
    STYLE_HEADING_2,
    STYLE_HEADING_3,
    STYLE_NORMAL,
    STYLE_NUM_NORMAL,
    STYLE_QUOTATION,
    apply_template_styles,
    cm_to_pt,
    indent_level_to_pt,
    resolve_style_name,
)
from transcribe_api.documents.word_document import WordDocument
from transcribe_api.runtime.logger import logger
from transcribe_api.documents.lib.markdown_to_blocks import markdown_to_plain_text
from transcribe_api.documents.lib.models import (
    Heading,
    HearingDocument,
    PageBreak,
    Paragraph,
    TemplateBody,
)
from transcribe_api.documents.lib.procedural_frontmatter_docx import render_frontmatter

if TYPE_CHECKING:
    from transcribe_api.api.routes_dictation import LiveTranscriptionSubmissionRequest

# Current year for copyright notice
CURRENT_YEAR = datetime.now(UTC).year
HEADING_LEVEL_2 = 2
MIN_QUOTE_INDENT_LEVEL = 2

# Template paths
# MERGE NOTE: was three parents up (backend/) under the old layout, which
# resolves to src/ here. data/ ships inside the package so it is present in a
# wheel/container too.
BASE_DIR = Path(__file__).resolve().parent.parent
TEMPLATE_DIR = BASE_DIR / "data"
DECISION_TEMPLATE_PATH = TEMPLATE_DIR / "decision-template.docx"
ENGLISH_CREST_PATH = TEMPLATE_DIR / "english-judicial-decision-crest.png"
SCOTTISH_CREST_PATH = TEMPLATE_DIR / "scottish-judicial-decision-crest.png"

# Locations that use the Scottish crest
SCOTTISH_LOCATIONS: set[str] = {"Glasgow"}


def _generate_output_filename(user_email: str) -> str:
    """Generate a unique filename for the output document."""
    timestamp = datetime.now(UTC).strftime("%Y%m%d_%H%M%S")
    safe_email = user_email.replace("@", "_at_").replace(".", "_")
    return f"hearing_document_{safe_email}_{timestamp}.docx"


# Controls that should be deleted entirely if no data is provided
# (rather than showing "N/A")
DELETE_IF_EMPTY_CONTROLS: set[str] = {
    "AnonymityHead",
}

# Anonymity-related content control keys (must match Word template tags)
ANONYMITY_CONTENT_CONTROLS = {
    "AnonymityHead",
    "AnonymityReasonsHeading",
    "AnonymityReasonsText",
    "AnonymityOrder",
    "AnonymityOrderText",
}

# Controls that should be left unchanged if no data provided
# (keeps original template text - useful for section headings)
SKIP_IF_EMPTY_CONTROLS: set[str] = {
    "RepresentationHeading",  # Keep "Representation" heading text
}

# Suffixes that indicate a control should be left unchanged if no data
# (keeps original template text)
SKIP_IF_EMPTY_SUFFIXES = ("Heading",)

# Section placeholder text in the Word template that will be replaced
# with transcript content via find-and-replace
SECTION_PLACEHOLDERS = {
    "background": "[ADD ANY RELEVANT BACKGROUND TO THE CASE]",
    "evidence": "[ORAL OR WRITTEN EVIDENCE THAT WAS CONSIDERED]",
    "facts": "[WRITE THIS]",
}

FEE_AWARD_PLACEHOLDER = "[Fee award to be determined]"
FEE_AWARD_DROPDOWN_PLACEHOLDER = "Select fee award"
FEE_AWARD_DROPDOWN_OPTIONS: tuple[str, ...] = (
    (
        "As I have allowed the appeal I have considered whether or not to make a fee award. I have decided to make a fee award for any fee which has been paid or is payable because there is no reason for the fee award not to follow the outcome of the appeal."
    ),
    (
        "As I have allowed the appeal I have considered whether or not to make a fee award. I have decided to make no fee award because [REASONS]."
    ),
    "I have dismissed the appeal and therefore there can be no fee award.",
    "No fee is paid or payable and therefore there can be no fee award.",
)

_ALIGNMENT_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


@dataclass(slots=True)
class _RenderState:
    in_decision_body: bool = False


def _apply_alignment(paragraph, alignment: str | None) -> None:
    if alignment:
        paragraph.alignment = _ALIGNMENT_MAP[alignment]


def _derive_indent_level(block: Paragraph) -> int:
    if block.left_indent_pt is None:
        return 0
    return max(round(block.left_indent_pt / indent_level_to_pt(1)), 0)


def _set_paragraph_list_level(paragraph, ilvl: int, num_id: int = 1) -> None:
    """Override the numbering level on a single paragraph via ``w:numPr``."""
    ppr = paragraph._element.get_or_add_pPr()  # noqa: SLF001
    existing = ppr.find(qn("w:numPr"))
    if existing is not None:
        ppr.remove(existing)
    numpr = OxmlElement("w:numPr")
    ilvl_elem = OxmlElement("w:ilvl")
    ilvl_elem.set(qn("w:val"), str(ilvl))
    numpr.append(ilvl_elem)
    numid_elem = OxmlElement("w:numId")
    numid_elem.set(qn("w:val"), str(num_id))
    numpr.append(numid_elem)
    ppr.append(numpr)


def _append_fee_award_dropdown(paragraph) -> None:
    sdt = OxmlElement("w:sdt")
    sdt_pr = OxmlElement("w:sdtPr")

    alias = OxmlElement("w:alias")
    alias.set(qn("w:val"), "FeeAwardSelection")
    sdt_pr.append(alias)

    tag = OxmlElement("w:tag")
    tag.set(qn("w:val"), "FeeAwardSelection")
    sdt_pr.append(tag)

    dropdown = OxmlElement("w:dropDownList")
    for option in FEE_AWARD_DROPDOWN_OPTIONS:
        list_item = OxmlElement("w:listItem")
        list_item.set(qn("w:displayText"), option)
        list_item.set(qn("w:value"), option)
        dropdown.append(list_item)
    sdt_pr.append(dropdown)
    sdt.append(sdt_pr)

    sdt_content = OxmlElement("w:sdtContent")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = FEE_AWARD_DROPDOWN_PLACEHOLDER
    run.append(text)
    sdt_content.append(run)
    sdt.append(sdt_content)

    paragraph._element.append(sdt)  # noqa: SLF001


def _apply_paragraph_indentation(
    paragraph,
    block: Paragraph,
    resolved_style: str,
    indent_level: int,
) -> None:
    if resolved_style == STYLE_NUM_NORMAL:
        # Word's numbering definition handles indentation per-level.
        # Only override when the block targets a sub-level (ilvl > 0).
        if block.list_level is not None and block.list_level > 0:
            _set_paragraph_list_level(paragraph, block.list_level)
        return
    if block.left_indent_pt is not None:
        if resolved_style == STYLE_QUOTATION:
            extra_levels = max(indent_level - MIN_QUOTE_INDENT_LEVEL, 0)
            paragraph.paragraph_format.left_indent = Pt(cm_to_pt(QUOTATION_BASE_INDENT_CM) + extra_levels * cm_to_pt(INDENT_INCREMENT_CM))
        else:
            paragraph.paragraph_format.left_indent = Pt(indent_level_to_pt(indent_level))
    if block.first_line_indent_pt is not None:
        paragraph.paragraph_format.first_line_indent = Pt(block.first_line_indent_pt)


def _append_paragraph_block(
    container,
    block: Paragraph,
    state: _RenderState,
    document,
    section_name: str,
    *,
    apply_decision_rules: bool,
) -> None:
    indent_level = _derive_indent_level(block)
    resolved_style = block.style
    if apply_decision_rules and state.in_decision_body and resolved_style is None:
        if section_name in ("signature", "fee_award"):
            resolved_style = STYLE_NORMAL
        else:
            resolved_style = (
                STYLE_QUOTATION if block.italic and indent_level >= MIN_QUOTE_INDENT_LEVEL else STYLE_NUM_NORMAL
            )

    resolved_style = resolve_style_name(document, resolved_style, fallback=STYLE_NORMAL)
    paragraph = container.add_paragraph(style=resolved_style)
    _apply_alignment(paragraph, block.alignment)
    if block.right_tab:
        section = paragraph.part.document.sections[0]
        usable_width = section.page_width - section.left_margin - section.right_margin
        paragraph.paragraph_format.tab_stops.add_tab_stop(usable_width, WD_TAB_ALIGNMENT.RIGHT)

    _apply_paragraph_indentation(paragraph, block, resolved_style, indent_level)

    if section_name == "fee_award" and block.text.strip() == FEE_AWARD_PLACEHOLDER:
        _append_fee_award_dropdown(paragraph)
        return

    run = paragraph.add_run(block.text)
    run.bold = block.bold
    run.italic = block.italic
    run.underline = block.underline
    run.font.color.rgb = None
    if block.font_size:
        run.font.size = Pt(block.font_size)


def _heading_style_name(level: int) -> str:
    if level <= 1:
        return STYLE_HEADING_1
    if level == HEADING_LEVEL_2:
        return STYLE_HEADING_2
    return STYLE_HEADING_3


def _append_heading_block(
    container,
    block: Heading,
    state: _RenderState,
    document,
    *,
    apply_decision_rules: bool,
) -> None:
    style_name = resolve_style_name(document, _heading_style_name(block.level), fallback=STYLE_HEADING_1)
    paragraph = container.add_paragraph(style=style_name)
    _apply_alignment(paragraph, block.alignment)
    run = paragraph.add_run(block.text)
    run.font.color.rgb = None
    run.font.size = Pt(block.font_size)
    is_decision_and_reasons = block.text.strip().upper() == "DECISION AND REASONS"
    if is_decision_and_reasons:
        run.bold = True
        run.underline = True
    if apply_decision_rules and is_decision_and_reasons:
        state.in_decision_body = True


def _add_page_number(paragraph) -> None:
    """Insert a PAGE field into a paragraph for automatic page numbering."""
    begin_run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    begin_run._r.append(fld_char_begin)  # noqa: SLF001

    instr_run = paragraph.add_run()
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    instr_run._r.append(instr_text)  # noqa: SLF001

    end_run = paragraph.add_run()
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    end_run._r.append(fld_char_end)  # noqa: SLF001


def _configure_after_first_page_header(root, appeal_number: str) -> None:
    section = root.sections[0]
    section.different_first_page_header_footer = True

    # Leave first page header empty.
    first_page_header = section.first_page_header
    first_para = first_page_header.paragraphs[0] if first_page_header.paragraphs else first_page_header.add_paragraph()
    first_para.text = ""

    # Set header for pages after the first page.
    header = section.header
    header_para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    header_para.text = f"Appeal number: {appeal_number}"
    header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header_para.runs:
        run.font.color.rgb = None

    # Set subsequent pages footer to a centred page number.
    footer = section.footer
    footer_para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    footer_para.clear()
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_page_number(footer_para)


def _render_structured_blocks(
    container,
    blocks: list,
    section_name: str,
    state: _RenderState,
    document,
    *,
    apply_decision_rules: bool = True,
) -> None:
    for block in blocks:
        if isinstance(block, Paragraph):
            _append_paragraph_block(
                container,
                block,
                state,
                document,
                section_name,
                apply_decision_rules=apply_decision_rules,
            )
        elif isinstance(block, Heading):
            _append_heading_block(
                container,
                block,
                state,
                document,
                apply_decision_rules=apply_decision_rules,
            )
        elif isinstance(block, PageBreak):
            # Header/footer containers do not support page breaks.
            if hasattr(container, "add_page_break"):
                container.add_page_break()
            else:
                logger.warning("Skipping page break in %s section", section_name)


def _render_template_body(root, template_body: TemplateBody, state: _RenderState) -> None:
    for section_name, blocks in template_body.ordered_sections():
        _render_structured_blocks(root, blocks, section_name, state, root, apply_decision_rules=True)


def render_structured_document(hearing_document: HearingDocument, output_path: Path) -> Path:
    """Render a .docx from a structured HearingDocument model."""
    doc = WordDocument()
    root = doc.document
    apply_template_styles(root)
    state = _RenderState()

    _render_structured_blocks(
        root.sections[0].header,
        hearing_document.header,
        "header",
        state,
        root,
        apply_decision_rules=False,
    )
    if hearing_document.frontmatter is not None:
        render_frontmatter(root, hearing_document.frontmatter, hearing_document.frontmatter.crest_path)
        appeal_number = hearing_document.frontmatter.case_id or hearing_document.frontmatter.appeal_number
        _configure_after_first_page_header(root, appeal_number)
    if isinstance(hearing_document.body, TemplateBody):
        _render_template_body(root, hearing_document.body, state)
    else:
        _render_structured_blocks(root, hearing_document.body, "body", state, root, apply_decision_rules=True)
    _render_structured_blocks(
        root.sections[0].first_page_footer,
        hearing_document.footer,
        "footer",
        state,
        root,
        apply_decision_rules=False,
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path


def render_document(  # noqa: C901, PLR0912, PLR0915
    document_data: DocumentData | HearingDocument,
    output_path: Path,
    template_path: Path | None = None,
    *,
    default_for_missing: str | None = "N/A",
    style_missing_as_placeholder: bool = False,
    delete_if_empty: set[str] | None = None,
) -> Path:
    """Render a document from DocumentData.

    Populates the template's content controls with form data and transcript.
    For more control (text replacement, appending content), use WordDocument directly.

    Args:
        document_data: Form data and transcript messages
        output_path: Where to save the rendered document
        template_path: Custom template path (defaults to decision template)
        default_for_missing: Value to use for controls without data (None to leave empty)
        style_missing_as_placeholder: If True, style missing values as bold red
        delete_if_empty: Controls to delete if no data (defaults to DELETE_IF_EMPTY_CONTROLS)

    Returns:
        Path to the rendered document

    Raises:
        FileNotFoundError: If template doesn't exist
        RuntimeError: If document generation fails
    """
    if isinstance(document_data, HearingDocument):
        return render_structured_document(document_data, output_path)

    template = template_path or DECISION_TEMPLATE_PATH

    if not template.exists():
        msg = f"Template not found: {template}"
        raise FileNotFoundError(msg)

    try:
        doc = WordDocument.from_template(template)
        apply_template_styles(doc.document)
        template_data = document_data.to_template_dict()

        # Look up actual content from IDs
        form_data = document_data.form_data

        # Replace hearing type ID with actual description
        if form_data.hearing_type:
            hearing_description = get_hearing_description(form_data.hearing_type)

            # Append Rule 28 text for appellant non-attendance
            if form_data.requires_appellant_rule_28_text():
                appellant_rule_28_text = (
                    "\n\nThe Appellant failed to attend the hearing. "
                    "I considered rule 28 of the Tribunal Procedure (First-tier Tribunal) "
                    "(Immigration and Asylum Chamber) Rules 2014. "
                    "I was satisfied that [the Appellant had actual notice of the hearing / "
                    "reasonable steps had been taken to notify the Appellant of the hearing] "
                    "because [REASONS]. It was in the interests of justice to proceed "
                    "without the Appellant because [REASONS]."
                )
                hearing_description += appellant_rule_28_text
                logger.info("Added Rule 28 text for appellant non-attendance")

            # Append Rule 28 text for respondent non-representation
            if form_data.requires_respondent_rule_28_text():
                respondent_rule_28_text = (
                    "\n\nThe Respondent was not represented at the hearing. "
                    "I considered rule 28 of the Tribunal Procedure (First-tier Tribunal) "
                    "(Immigration and Asylum Chamber) Rules 2014. "
                    "I was satisfied that [the Respondent had actual notice of the hearing / "
                    "reasonable steps had been taken to notify the Respondent of the hearing] "
                    "because [REASONS]. It was in the interests of justice to proceed "
                    "without the Respondent because [REASONS]."
                )
                hearing_description += respondent_rule_28_text
                logger.info("Added Rule 28 text for respondent non-representation")

            template_data["HearingDescription"] = markdown_to_plain_text(hearing_description)
            template_data["HearingTitle"] = get_hearing_title(form_data.hearing_type)

        # Replace legal issues IDs with display titles and content
        if form_data.legal_issues:
            template_data["Issues"] = get_legal_issues_display(form_data.legal_issues)
            template_data["LegalFramework"] = get_legal_framework_content(form_data.legal_issues)

        # FeeAward - default placeholder (can be customized based on form data)
        if "FeeAward" not in template_data or not template_data.get("FeeAward"):
            template_data["FeeAward"] = FEE_AWARD_PLACEHOLDER

        # Note: CaseID is used both in the body and in page headers (same tag name)
        # The populate_content_controls method now searches headers, so both get populated

        # Compute controls to delete or skip
        template_controls = doc.get_content_control_keys()

        controls_to_delete = delete_if_empty if delete_if_empty is not None else DELETE_IF_EMPTY_CONTROLS.copy()

        # Process anonymity content based on the anonymity order status
        anonymity_status = parse_anonymity_status(form_data.anonymity_order)
        anonymity_content = get_anonymity_content(anonymity_status)

        # Merge anonymity data into template data
        template_data.update(anonymity_content.to_template_dict())

        # Add anonymity controls marked for deletion
        for key, value in anonymity_content.to_template_dict().items():
            if value == DELETE_MARKER:
                controls_to_delete.add(key)

        logger.info(
            "Anonymity status: %s, controls to delete: %s",
            anonymity_status.value,
            anonymity_content.get_delete_controls(),
        )

        # Compute controls to skip (leave unchanged)
        controls_to_skip = SKIP_IF_EMPTY_CONTROLS.copy()
        for key in template_controls:
            if key.endswith(SKIP_IF_EMPTY_SUFFIXES):
                controls_to_skip.add(key)

        logger.info(
            "Template data keys: %s",
            sorted(template_data.keys()),
        )

        result = doc.populate_content_controls(
            template_data,
            default_for_missing=default_for_missing,
            style_missing_as_placeholder=style_missing_as_placeholder,
            delete_if_empty=controls_to_delete,
            skip_if_empty=controls_to_skip,
        )

        logger.info(
            "Document populated: %d fields filled, %d missing, %d unused keys",
            len(result.populated_fields),
            len(result.missing_fields),
            len(result.unused_data_keys),
        )

        if result.missing_fields:
            logger.warning("Missing template fields: %s", sorted(result.missing_fields))

        # Replace section placeholders with transcript content
        section_transcripts = document_data.get_section_transcripts()
        sections_replaced = []

        for section_name, placeholder in SECTION_PLACEHOLDERS.items():
            section_content = section_transcripts.format_section(section_name)
            if section_content:
                replacements = doc.find_and_replace(placeholder, section_content)
                if replacements > 0:
                    sections_replaced.append(section_name)
                    logger.info(
                        "Replaced %d instance(s) of '%s' placeholder with %d characters of %s transcript",
                        replacements,
                        section_name,
                        len(section_content),
                        section_name,
                    )

        if sections_replaced:
            logger.info("Section transcripts inserted: %s", sections_replaced)

        # Update copyright year in footer (replace any year with current year)
        copyright_replacements = doc.find_and_replace(
            r"CROWN COPYRIGHT \d{4}",
            f"CROWN COPYRIGHT {CURRENT_YEAR}",
            regex=True,
        )
        if copyright_replacements > 0:
            logger.info(
                "Updated %d copyright notice(s) to year %d",
                copyright_replacements,
                CURRENT_YEAR,
            )

        # Swap crest to Scottish version for Glasgow hearings
        # (the template already contains the English crest by default)
        if form_data.location in SCOTTISH_LOCATIONS:
            if doc.replace_crest_image(SCOTTISH_CREST_PATH):
                logger.info("Replaced crest image with Scottish crest")
            else:
                logger.warning("No crest image found to replace")

        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(output_path)

    except Exception as e:
        logger.exception("Document generation failed")
        msg = f"Document generation failed: {e}"
        raise RuntimeError(msg) from e

    return output_path


def render_hearing_document(
    request: LiveTranscriptionSubmissionRequest,
    output_dir: Path,
    user_email: str,
) -> Path:
    """Render a hearing document from an API submission request.

    This is the interface used by the /live-transcription/submission endpoint.
    It routes through the structured HearingDocument pipeline by default.

    Args:
        request: The live transcription submission request
        output_dir: Directory for the rendered document
        user_email: Email of the submitting user

    Returns:
        Path to the rendered document
    """
    return render_hearing_document_structured(request, output_dir, user_email)


def render_hearing_document_structured(
    request: LiveTranscriptionSubmissionRequest,
    output_dir: Path,
    user_email: str,
) -> Path:
    """Render a hearing document through HearingDocument block routing."""
    document_data = _build_document_data_from_submission(request, user_email)
    hearing_document = build_hearing_document_from_document_data(document_data)

    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / _generate_output_filename(user_email)
    return render_document(hearing_document, output_path)


def _build_document_data_from_submission(
    request: LiveTranscriptionSubmissionRequest,
    user_email: str,
) -> DocumentData:
    from transcribe_api.documents.models import SectionTranscripts

    form = request.form_data

    # Get sectioned messages from the request (handles both sectioned and legacy formats)
    sectioned_msgs = request.get_sectioned_messages()

    # Convert API messages to document model messages
    def convert_messages(msgs: list) -> list[TranscriptMessage]:
        return [
            TranscriptMessage(
                speaker=msg.speaker,
                text=msg.text,
                timestamp=msg.timestamp,
                timestamp_ms=msg.timestamp_ms,
            )
            for msg in msgs
        ]

    # Create sectioned transcript for document data
    section_transcripts = SectionTranscripts(
        background=convert_messages(sectioned_msgs.background),
        evidence=convert_messages(sectioned_msgs.evidence),
        facts=convert_messages(sectioned_msgs.facts),
    )

    document_data = DocumentData(
        notes=getattr(request, "notes", None),
        form_data=HearingFormData(
            case_id=form.case_id,
            location=form.location,
            location_other=form.location_other,
            jurisdiction=form.jurisdiction,
            hearing_date=form.hearing_date,
            judge_name=form.judge_name,
            anonymity_order=form.anonymity_order,
            appellant_name=form.appellant_name,
            respondent=form.respondent,
            hearing_type=form.hearing_type,
            # Appellant representation fields
            appellant_rep_type=form.appellant_rep_type,
            appellant_rep_details=form.appellant_rep_details,
            # Respondent representation fields
            respondent_rep_type=form.respondent_rep_type,
            respondent_rep_name=form.respondent_rep_name,
            appealable_decision_date=form.appealable_decision_date,
            legal_issues=form.legal_issues,
            document_type=form.document_type,
            next_hearing_type=form.next_hearing_type,
            next_hearing_adjudicator=form.next_hearing_adjudicator,
        ),
        messages=section_transcripts,
        user_email=user_email,
    )
    return document_data
