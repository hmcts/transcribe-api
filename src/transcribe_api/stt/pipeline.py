import asyncio
import tempfile
from datetime import UTC, datetime
from pathlib import Path
from uuid import UUID

import sentry_sdk
from docx import Document as DocxDocument

from transcribe_api.runtime.blob_dictation import AsyncAzureBlobManager
from transcribe_api.stt.speakers_dictation import process_speakers_and_dialogue_entries
from transcribe_api.stt.fast_client import transcribe_audio
from transcribe_api.stt.utils import (
    get_file_blob_path,
    get_url_for_transcription,
)
from transcribe_api.domain.interface_dictation import (
    save_transcription,
    save_transcription_job,
)
from transcribe_api.domain.models_dictation import (
    Transcription,
    TranscriptionJob,
)
from transcribe_api.runtime.logger import logger
from transcribe_api.documents.minutes.llm_calls import (
    generate_llm_output_task,
    generate_meeting_title,
)
from transcribe_api.documents.minutes.templates.templates_metadata import (
    crissa_template,
    general_template,
)
from transcribe_api.domain.notify.gov_notify import send_email


async def generate_and_save_meeting_title(
    dialogue_entries: list, transcription: Transcription, user_id: UUID, user_email: str
) -> Transcription:
    try:
        logger.info(f"Generating meeting title for {len(dialogue_entries)} dialogue entries")
        provisional_title = await generate_meeting_title(dialogue_entries, user_email)
        existing_transcription = transcription
        existing_transcription.title = provisional_title
        return save_transcription(existing_transcription, user_id)
    except Exception as e:
        logger.error(f"Error saving transcription: {e}")
        sentry_sdk.capture_exception(e)
        return transcription


async def _generate_and_upload_transcript_docx(
    dialogue_entries: list,
    azure_user_id: str,
    title: str,
) -> str | None:
    """Generate a plain .docx transcript from dialogue entries and upload to blob storage.

    Returns the blob path on success, None on failure.
    """
    try:
        timestamp = datetime.now(UTC).strftime("%Y%m%d_%H%M%S")
        doc_filename = f"transcript-{timestamp}.docx"
        doc_blob_path = get_file_blob_path(azure_user_id, doc_filename)

        output_dir = Path(tempfile.gettempdir()) / "transcripts"
        output_dir.mkdir(parents=True, exist_ok=True)
        temp_path = output_dir / doc_filename

        doc = DocxDocument()
        doc.add_heading(title, level=0)

        for entry in dialogue_entries:
            speaker = entry.speaker if hasattr(entry, "speaker") else entry.get("speaker", "Unknown")
            text = entry.text if hasattr(entry, "text") else entry.get("text", "")
            para = doc.add_paragraph()
            run = para.add_run(f"{speaker}: ")
            run.bold = True
            para.add_run(text)

        doc.save(str(temp_path))

        async with AsyncAzureBlobManager() as blob_manager:
            success = await blob_manager.create_blob_from_file(temp_path, doc_blob_path)

        if success:
            logger.info("Uploaded transcript document to blob %s for user %s", doc_blob_path, azure_user_id)
            return doc_blob_path
        else:
            logger.warning("Failed to upload transcript document for user %s", azure_user_id)
            return None

    except Exception as e:
        logger.error("Error generating transcript document: %s", e)
        sentry_sdk.capture_exception(e)
        return None


async def transcribe_and_generate_llm_output(
    user_upload_blob_storage_file_key: str,
    user_id: UUID,
    user_email: str,
    azure_user_id: str,
    transcription_id: str | None = None,
):
    # Start a Sentry transaction for the whole function
    with sentry_sdk.start_transaction(op="task", name="Transcribe and Generate LLM Output") as transaction:  # noqa: F841
        transcription_data = Transcription(id=transcription_id)
        transcription = save_transcription(transcription_data, user_id)

        try:
            dialogue_entries = await transcribe_audio(user_upload_blob_storage_file_key)
            updated_dialogue_entries = await process_speakers_and_dialogue_entries(dialogue_entries, user_email)
            saved_job = save_transcription_job(
                TranscriptionJob(
                    transcription_id=transcription.id,
                    dialogue_entries=updated_dialogue_entries,
                    s3_audio_url=user_upload_blob_storage_file_key,
                )
            )

        except Exception as e:
            save_transcription_job(
                TranscriptionJob(
                    transcription_id=transcription.id,
                    dialogue_entries=[],
                    s3_audio_url=user_upload_blob_storage_file_key,
                    error_message=str(e),
                )
            )
            sentry_sdk.capture_exception(e)
            raise

        # Start all three tasks in parallel
        general_task = generate_llm_output_task(
            updated_dialogue_entries, transcription.id, general_template, user_email
        )
        title_task = generate_and_save_meeting_title(updated_dialogue_entries, transcription, user_id, user_email)
        crissa_task = generate_llm_output_task(updated_dialogue_entries, transcription.id, crissa_template, user_email)

        try:
            # Wait for all three tasks including CRISSA before sending email
            await asyncio.gather(general_task, title_task, crissa_task)
        except Exception as e:
            logger.error(f"Error in parallel tasks: {e}")
            sentry_sdk.capture_exception(e)

        # Generate a .docx transcript document and update s3_audio_url so that
        # "Download Transcript" serves the document rather than the raw audio file.
        # title_task mutates transcription.title so it is available here.
        doc_blob_path = await _generate_and_upload_transcript_docx(
            updated_dialogue_entries,
            azure_user_id,
            transcription.title or "Transcript",
        )
        if doc_blob_path:
            saved_job.s3_audio_url = doc_blob_path
            save_transcription_job(saved_job)

        try:
            send_email(
                user_email,
                get_url_for_transcription(transcription.id),
                transcription.title or "",
            )
        except Exception as e:
            logger.error(f"Error sending email: {e}")
            sentry_sdk.capture_exception(e)
