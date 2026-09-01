"""Unit tests for structured HearingDocument rendering."""

from math import isclose
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

from defusedxml.ElementTree import fromstring
from docx import Document
from docx.oxml.ns import qn

from transcribe_api.documents.styles import QUOTATION_BASE_INDENT_CM, cm_to_pt
from transcribe_api.documents.template_renderer import (
    FEE_AWARD_DROPDOWN_OPTIONS,
    FEE_AWARD_DROPDOWN_PLACEHOLDER,
    render_document,
    render_hearing_document,
    render_structured_document,
)
from transcribe_api.documents.lib.models import (
    Frontmatter,
    Heading,
    HearingDocument,
    IACTemplate,
    PageBreak,
    Paragraph,
)


def test_render_structured_document_writes_header_body_footer(tmp_path: Path) -> None:
    frontmatter = Frontmatter(
        case_id="PA/12345/2025",
        appeal_number="TST-123",
        hearing_location="Manchester",
        hearing_date="23 January 2026",
        judge_name="N SONI",
        appellant_name="ELLIOTT",
        appellant_representative="Ms. Patel",
        respondent_representative="Mr. Thomas",
    )
    hearing_document = HearingDocument(
        header=[Paragraph(text="Header Case Ref", bold=True)],
        frontmatter=frontmatter,
        body=[
            Heading(text="Section Heading"),
            Paragraph(text="Body paragraph one"),
            PageBreak(),
            Paragraph(text="Body paragraph two"),
        ],
        footer=[Paragraph(text="Footer text")],
    )
    output_path = tmp_path / "structured.docx"

    render_structured_document(hearing_document, output_path)

    generated = Document(str(output_path))
    body_text = [p.text for p in generated.paragraphs if p.text.strip()]
    assert "THE IMMIGRATION ACTS" in body_text
    assert "Representation:" in body_text
    assert "Section Heading" in body_text
    assert "Body paragraph one" in body_text
    assert "Body paragraph two" in body_text
    heading_para = next(paragraph for paragraph in generated.paragraphs if paragraph.text == "Section Heading")
    assert heading_para.style.name == "Heading 1"
    assert heading_para.style.font.bold is True

    header_text = [p.text for p in generated.sections[0].header.paragraphs if p.text.strip()]
    assert "Header Case Ref" in header_text

    first_page_footer_text = [p.text for p in generated.sections[0].first_page_footer.paragraphs if p.text.strip()]
    assert "Footer text" in first_page_footer_text

    footer_element = generated.sections[0].footer._element
    instr_texts = footer_element.findall(f".//{qn('w:instrText')}")
    assert any("PAGE" in (elem.text or "") for elem in instr_texts)


def test_render_structured_document_applies_global_style_defaults(tmp_path: Path) -> None:
    frontmatter = Frontmatter(
        appeal_number="TST-123",
        hearing_location="Manchester",
        hearing_date="23 January 2026",
        judge_name="N SONI",
        appellant_name="ELLIOTT",
    )
    hearing_document = HearingDocument(
        frontmatter=frontmatter,
        body=[Heading(text="DECISION AND REASONS", level=1), Paragraph(text="A first numbered paragraph")],
    )
    output_path = tmp_path / "formatted.docx"

    render_structured_document(hearing_document, output_path)

    generated = Document(str(output_path))
    section = generated.sections[0]
    assert isclose(section.top_margin.cm, 2.54, abs_tol=0.01)
    assert isclose(section.right_margin.cm, 2.0, abs_tol=0.01)
    assert isclose(section.bottom_margin.cm, 2.54, abs_tol=0.01)
    assert isclose(section.left_margin.cm, 2.0, abs_tol=0.01)

    normal = generated.styles["Normal"]
    assert normal.font.name == "Arial"
    assert isclose(normal.font.size.pt, 12.0, abs_tol=0.01)

    h1 = generated.styles["Heading 1"]
    h2 = generated.styles["Heading 2"]
    h3 = generated.styles["Heading 3"]
    assert h1.font.bold is True
    assert h2.font.underline is True
    assert h3.font.italic is True


def test_render_structured_document_formats_decision_and_reasons_heading(tmp_path: Path) -> None:
    frontmatter = Frontmatter(
        appeal_number="TST-123",
        hearing_location="Manchester",
        hearing_date="23 January 2026",
        judge_name="N SONI",
        appellant_name="ELLIOTT",
    )
    hearing_document = HearingDocument(
        frontmatter=frontmatter,
        body=[Heading(text="DECISION AND REASONS", level=1), Paragraph(text="A first numbered paragraph")],
    )
    output_path = tmp_path / "decision_heading.docx"

    render_structured_document(hearing_document, output_path)

    generated = Document(str(output_path))
    decision_heading = next(paragraph for paragraph in generated.paragraphs if paragraph.text == "DECISION AND REASONS")
    non_empty_runs = [run for run in decision_heading.runs if run.text.strip()]
    assert non_empty_runs
    assert any(run.bold is True for run in non_empty_runs)
    assert any(run.underline is True for run in non_empty_runs)


def test_render_structured_document_inserts_fee_award_dropdown(tmp_path: Path) -> None:
    frontmatter = Frontmatter(
        appeal_number="TST-123",
        hearing_location="Manchester",
        hearing_date="23 January 2026",
        judge_name="N SONI",
        appellant_name="ELLIOTT",
    )
    hearing_document = HearingDocument(
        frontmatter=frontmatter,
        body=IACTemplate(
            fee_award=[
                Heading(text="Fee Award", level=2, bold=True, underline=True),
                Paragraph(text="[Fee award to be determined]", alignment="justify"),
            ]
        ),
    )
    output_path = tmp_path / "fee_award_dropdown.docx"

    render_structured_document(hearing_document, output_path)

    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with ZipFile(output_path) as docx_zip:
        document_root = fromstring(docx_zip.read("word/document.xml"))

    dropdown = document_root.find(".//w:dropDownList", ns)
    assert dropdown is not None

    list_items = dropdown.findall("./w:listItem", ns)
    values = [item.get(qn("w:value")) for item in list_items]
    assert values == list(FEE_AWARD_DROPDOWN_OPTIONS)

    all_text = "".join(node.text or "" for node in document_root.findall(".//w:t", ns))
    assert "[Fee award to be determined]" not in all_text
    assert FEE_AWARD_DROPDOWN_PLACEHOLDER in all_text


def test_render_structured_document_marks_headings_as_word_headings(tmp_path: Path) -> None:
    frontmatter = Frontmatter(
        appeal_number="TST-123",
        hearing_location="Manchester",
        hearing_date="23 January 2026",
        judge_name="N SONI",
        appellant_name="ELLIOTT",
    )
    hearing_document = HearingDocument(
        frontmatter=frontmatter,
        body=[
            Heading(text="Top Level Heading", level=1),
            Heading(text="Second Level Heading", level=2),
            Heading(text="Third Level Heading", level=3),
        ],
    )
    output_path = tmp_path / "heading_levels.docx"

    render_structured_document(hearing_document, output_path)

    generated = Document(str(output_path))
    paragraphs = {p.text: p for p in generated.paragraphs if p.text.strip()}
    assert paragraphs["Top Level Heading"].style.name == "Heading 1"
    assert paragraphs["Second Level Heading"].style.name == "Heading 2"
    assert paragraphs["Third Level Heading"].style.name == "Heading 3"

    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with ZipFile(output_path) as docx_zip:
        document_root = fromstring(docx_zip.read("word/document.xml"))
        styles_root = fromstring(docx_zip.read("word/styles.xml"))

    def paragraph_style_val(text: str) -> str | None:
        for paragraph in document_root.findall(".//w:body/w:p", ns):
            text_value = "".join(node.text or "" for node in paragraph.findall(".//w:t", ns))
            if text_value == text:
                p_style = paragraph.find("./w:pPr/w:pStyle", ns)
                return p_style.get(qn("w:val")) if p_style is not None else None
        return None

    def style_outline_val(style_id: str) -> str | None:
        for style in styles_root.findall("./w:style", ns):
            if style.get(qn("w:styleId")) == style_id:
                outline = style.find("./w:pPr/w:outlineLvl", ns)
                return outline.get(qn("w:val")) if outline is not None else None
        return None

    assert paragraph_style_val("Top Level Heading") == "Heading1"
    assert paragraph_style_val("Second Level Heading") == "Heading2"
    assert paragraph_style_val("Third Level Heading") == "Heading3"

    assert style_outline_val("Heading1") == "0"
    assert style_outline_val("Heading2") == "1"
    assert style_outline_val("Heading3") == "2"


def _has_word_numbering(paragraph) -> bool:
    """Return True when the paragraph carries a w:numPr element (style or direct)."""
    ppr = paragraph._element.find(qn("w:pPr"))
    if ppr is not None and ppr.find(qn("w:numPr")) is not None:
        return True
    style_elem = paragraph.style.element if paragraph.style else None
    if style_elem is not None:
        sty_ppr = style_elem.find(qn("w:pPr"))
        if sty_ppr is not None and sty_ppr.find(qn("w:numPr")) is not None:
            return True
    return False


def test_render_structured_document_numbers_top_level_after_decision_heading(tmp_path: Path) -> None:
    frontmatter = Frontmatter(
        appeal_number="TST-123",
        hearing_location="Manchester",
        hearing_date="23 January 2026",
        judge_name="N SONI",
        appellant_name="ELLIOTT",
    )
    hearing_document = HearingDocument(
        frontmatter=frontmatter,
        body=[
            Heading(text="DECISION AND REASONS", level=1),
            Paragraph(text="First body paragraph"),
            Paragraph(text="Second body paragraph"),
        ],
    )
    output_path = tmp_path / "numbered.docx"

    render_structured_document(hearing_document, output_path)

    generated = Document(str(output_path))
    paragraphs = [p for p in generated.paragraphs if p.text.strip()]

    num_normal = [p for p in paragraphs if p.style.name == "NumNormal"]
    assert any(p.text == "First body paragraph" for p in num_normal)
    assert any(p.text == "Second body paragraph" for p in num_normal)

    for para in num_normal:
        assert _has_word_numbering(para), f"Expected Word numbering on: {para.text}"


def test_render_structured_document_does_not_number_signature_section(tmp_path: Path) -> None:
    frontmatter = Frontmatter(
        appeal_number="TST-123",
        hearing_location="Manchester",
        hearing_date="23 January 2026",
        judge_name="N SONI",
        appellant_name="ELLIOTT",
    )
    hearing_document = HearingDocument(
        frontmatter=frontmatter,
        body=IACTemplate(
            anonymity=[Heading(text="DECISION AND REASONS", level=1)],
            signature=[
                Heading(text="Signature", level=2, bold=True),
                Paragraph(text="Signed\tDate [not final until dated]", right_tab=True),
                Paragraph(text="[ADD SIGNATURE]", italic=True),
                Paragraph(text="First-tier Tribunal Judge N SONI"),
            ],
        ),
    )
    output_path = tmp_path / "signature_not_numbered.docx"

    render_structured_document(hearing_document, output_path)

    generated = Document(str(output_path))
    signature_paragraphs = [
        p
        for p in generated.paragraphs
        if p.text in {"Signed\tDate [not final until dated]", "[ADD SIGNATURE]", "First-tier Tribunal Judge N SONI"}
    ]
    assert signature_paragraphs
    for para in signature_paragraphs:
        assert para.style.name != "NumNormal"
        assert not _has_word_numbering(para), f"Did not expect Word numbering on signature paragraph: {para.text}"


def test_render_structured_document_uses_black_default_text(tmp_path: Path) -> None:
    frontmatter = Frontmatter(
        appeal_number="TST-123",
        hearing_location="Manchester",
        hearing_date="23 January 2026",
        judge_name="N SONI",
        appellant_name="ELLIOTT",
    )
    hearing_document = HearingDocument(
        frontmatter=frontmatter,
        body=[Heading(text="DECISION AND REASONS", level=1), Paragraph(text="Colour baseline paragraph")],
    )
    output_path = tmp_path / "colors.docx"
    render_structured_document(hearing_document, output_path)

    generated = Document(str(output_path))
    normal = generated.styles["Normal"]
    assert normal.font.color.rgb is None


def test_render_document_accepts_hearing_document_model(tmp_path: Path) -> None:
    frontmatter = Frontmatter(
        appeal_number="TST-123",
        hearing_location="Manchester",
        hearing_date="23 January 2026",
        judge_name="N SONI",
        appellant_name="ELLIOTT",
    )
    hearing_document = HearingDocument(
        frontmatter=frontmatter,
        body=[Paragraph(text="Body via render_document")],
    )
    output_path = tmp_path / "render_document_structured.docx"

    render_document(hearing_document, output_path)

    generated = Document(str(output_path))
    body_text = [p.text for p in generated.paragraphs if p.text.strip()]
    assert "THE IMMIGRATION ACTS" in body_text
    assert "Body via render_document" in body_text
    assert generated.sections[0].different_first_page_header_footer is True

    header_text = [p.text for p in generated.sections[0].header.paragraphs if p.text.strip()]
    assert "Appeal number: TST-123" in header_text


def test_render_structured_document_does_not_number_fee_award_section(tmp_path: Path) -> None:
    frontmatter = Frontmatter(
        appeal_number="TST-123",
        hearing_location="Manchester",
        hearing_date="23 January 2026",
        judge_name="N SONI",
        appellant_name="ELLIOTT",
    )
    hearing_document = HearingDocument(
        frontmatter=frontmatter,
        body=IACTemplate(
            anonymity=[Heading(text="DECISION AND REASONS", level=1)],
            fee_award=[
                Paragraph(text="Fee award paragraph one", alignment="justify"),
                Paragraph(text="Fee award paragraph two", alignment="justify"),
            ],
        ),
    )
    output_path = tmp_path / "fee_award_not_numbered.docx"

    render_structured_document(hearing_document, output_path)

    generated = Document(str(output_path))
    fee_award_paragraphs = [
        p
        for p in generated.paragraphs
        if p.text in {"Fee award paragraph one", "Fee award paragraph two"}
    ]
    assert fee_award_paragraphs
    for para in fee_award_paragraphs:
        assert para.style.name != "NumNormal"
        assert not _has_word_numbering(para), f"Did not expect numbering on fee award paragraph: {para.text}"


def test_render_structured_document_quotation_applies_2cm_base_indent(tmp_path: Path) -> None:
    # Minimum quote indent level is 2 (matches _MIN_QUOTE_INDENT_LEVEL in markdown_to_blocks).
    # At this level, the indent should be exactly the 2cm base with no additional increment.
    min_quote_indent_pt = cm_to_pt(2.0)  # INDENT_INCREMENT_PT * _MIN_QUOTE_INDENT_LEVEL
    frontmatter = Frontmatter(
        appeal_number="TST-123",
        hearing_location="Manchester",
        hearing_date="23 January 2026",
        judge_name="N SONI",
        appellant_name="ELLIOTT",
    )
    hearing_document = HearingDocument(
        frontmatter=frontmatter,
        body=IACTemplate(
            background=[
                Paragraph(
                    text="A quoted passage at the default indent.",
                    style="Quotation",
                    italic=True,
                    left_indent_pt=min_quote_indent_pt,
                )
            ]
        ),
    )
    output_path = tmp_path / "quotation_indent.docx"

    render_structured_document(hearing_document, output_path)

    generated = Document(str(output_path))
    quoted_para = next(p for p in generated.paragraphs if p.text == "A quoted passage at the default indent.")
    assert quoted_para.paragraph_format.left_indent is not None
    assert isclose(quoted_para.paragraph_format.left_indent.cm, QUOTATION_BASE_INDENT_CM, abs_tol=0.02)


def test_render_structured_document_quotation_increments_1cm_per_nesting_level(tmp_path: Path) -> None:
    min_quote_indent_pt = cm_to_pt(2.0)
    nested_quote_indent_pt = cm_to_pt(3.0)  # one level deeper
    frontmatter = Frontmatter(
        appeal_number="TST-123",
        hearing_location="Manchester",
        hearing_date="23 January 2026",
        judge_name="N SONI",
        appellant_name="ELLIOTT",
    )
    hearing_document = HearingDocument(
        frontmatter=frontmatter,
        body=IACTemplate(
            background=[
                Paragraph(
                    text="Default indent passage.",
                    style="Quotation",
                    italic=True,
                    left_indent_pt=min_quote_indent_pt,
                ),
                Paragraph(
                    text="Nested indent passage.",
                    style="Quotation",
                    italic=True,
                    left_indent_pt=nested_quote_indent_pt,
                ),
            ]
        ),
    )
    output_path = tmp_path / "quotation_nested_indent.docx"

    render_structured_document(hearing_document, output_path)

    generated = Document(str(output_path))
    paragraphs = {p.text: p for p in generated.paragraphs if p.text.strip()}
    default_para = paragraphs["Default indent passage."]
    nested_para = paragraphs["Nested indent passage."]

    assert isclose(default_para.paragraph_format.left_indent.cm, 2.0, abs_tol=0.02)
    assert isclose(nested_para.paragraph_format.left_indent.cm, 3.0, abs_tol=0.02)


def test_render_hearing_document_uses_structured_pipeline_without_template(monkeypatch, tmp_path: Path) -> None:
    # If old template pipeline were still used, this non-existent template path would fail.
    from transcribe_api.documents import template_renderer

    monkeypatch.setattr(template_renderer, "DECISION_TEMPLATE_PATH", tmp_path / "does-not-exist.docx")

    form_data = SimpleNamespace(
        case_id="PA/12345/2025",
        location="Taylor House",
        location_other=None,
        jurisdiction="First-tier Tribunal Immigration and Asylum Chamber",
        hearing_date="2026-01-23",
        judge_name="N SONI",
        anonymity_order="granted",
        appellant_name="ELLIOTT",
        respondent="THE SECRETARY OF STATE FOR THE HOME DEPARTMENT",
        hearing_type="face_to_face",
        appellant_rep_type="represented",
        appellant_rep_details="Ms. Patel",
        respondent_rep_type="counsel",
        respondent_rep_name="Mr. Thomas",
        appealable_decision_date="2025-12-01",
        legal_issues=["asylum_pre_naba", "credibility"],
        document_type="Decision",
        next_hearing_type=None,
        next_hearing_adjudicator=None,
    )
    sectioned_messages = SimpleNamespace(
        background=[SimpleNamespace(speaker="Judge", text="Background text", timestamp="10:00:00", timestamp_ms=0)],
        evidence=[SimpleNamespace(speaker="Counsel", text="Evidence text", timestamp="10:10:00", timestamp_ms=600000)],
        facts=[SimpleNamespace(speaker="Judge", text="Facts text", timestamp="10:20:00", timestamp_ms=1200000)],
    )
    request = SimpleNamespace(form_data=form_data, get_sectioned_messages=lambda: sectioned_messages)

    output_path = render_hearing_document(request, tmp_path, "court.clerk@justice.gov.uk")
    generated = Document(str(output_path))
    body_text = [p.text for p in generated.paragraphs if p.text.strip()]

    assert "The Hearing" in body_text
    assert "Background" in body_text
    assert "Evidence" in body_text
    assert "Facts" in body_text
