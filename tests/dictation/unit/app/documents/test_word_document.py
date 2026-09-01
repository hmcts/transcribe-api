"""Unit tests for low-level Word document editing behavior."""

from docx import Document
from docx.oxml import OxmlElement

from transcribe_api.documents.word_document import WordDocument


def test_find_and_replace_splits_double_newline_into_paragraphs() -> None:
    """Double newlines should create separate Word paragraphs."""
    doc = Document()
    doc.add_paragraph("[WRITE THIS]")
    word_doc = WordDocument(document=doc)

    replacements = word_doc.find_and_replace("[WRITE THIS]", "First para\n\nSecond para")

    assert replacements == 1
    assert len(doc.paragraphs) == 2
    assert doc.paragraphs[0].text == "First para"
    assert doc.paragraphs[1].text == "Second para"


def test_delete_content_control_removes_empty_parent_paragraph() -> None:
    """Deleting the only control in a paragraph should remove that paragraph."""
    doc = Document()
    paragraph = doc.add_paragraph()
    sdt = OxmlElement("w:sdt")
    paragraph._element.append(sdt)
    word_doc = WordDocument(document=doc)

    word_doc._delete_content_control(sdt)

    assert len(doc.paragraphs) == 0


def test_delete_content_control_keeps_non_empty_parent_paragraph() -> None:
    """A paragraph with visible text should not be removed."""
    doc = Document()
    paragraph = doc.add_paragraph("Keep me")
    sdt = OxmlElement("w:sdt")
    paragraph._element.append(sdt)
    word_doc = WordDocument(document=doc)

    word_doc._delete_content_control(sdt)

    assert len(doc.paragraphs) == 1
    assert doc.paragraphs[0].text == "Keep me"
