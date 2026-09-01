"""Unit tests for typography helpers."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from transcribe_api.documents.lib.typography import right_aligned_text


def test_right_aligned_text_sets_right_alignment() -> None:
    document = Document()
    right_aligned_text(document, "Appeal Number: TST-123")
    paragraph = document.paragraphs[-1]
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.RIGHT


def test_right_aligned_text_defaults_to_bold_no_underline() -> None:
    """Default call (appeal number) must be bold and not underlined."""
    document = Document()
    right_aligned_text(document, "Appeal Number: TST-123")
    run = document.paragraphs[-1].runs[0]
    assert run.bold is True
    assert run.underline is not True


def test_right_aligned_text_bold_false_underline_true() -> None:
    """Role-label call (Appellant/Respondent) must be underlined and not bold."""
    document = Document()
    right_aligned_text(document, "Appellant", bold=False, underline=True)
    run = document.paragraphs[-1].runs[0]
    assert run.bold is not True
    assert run.underline is True


def test_right_aligned_text_has_no_tab_character() -> None:
    """No tab character must appear in any run — absence of tabs is the fix for National Archives."""
    document = Document()
    right_aligned_text(document, "Respondent", bold=False, underline=True)
    assert not any("\t" in run.text for run in document.paragraphs[-1].runs)
