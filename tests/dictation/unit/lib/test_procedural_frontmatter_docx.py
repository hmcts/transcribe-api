"""Unit tests for procedural frontmatter DOCX generation."""

from math import isclose
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from transcribe_api.documents.lib.models import Frontmatter
from transcribe_api.documents.lib.procedural_frontmatter_docx import build_frontmatter_document


def _build_sample_frontmatter(*, include_anonymity_label: bool = True) -> Frontmatter:
    return Frontmatter(
        tribunal_title="First-tier Tribunal",
        tribunal_chamber="(Immigration and Asylum Chamber)",
        appeal_number="TST-123",
        acts_heading="THE IMMIGRATION ACTS",
        hearing_location_prefix="Heard at",
        hearing_location="Manchester",
        hearing_date_prefix="On",
        hearing_date="23 January 2026",
        before_heading="Before",
        judge_title="FIRST-TIER TRIBUNAL JUDGE",
        judge_name="N SONI",
        appellant_name="ELLIOTT",
        include_anonymity_label=include_anonymity_label,
    )


def _normalized_paragraph_texts(document_path: Path) -> list[str]:
    document = Document(str(document_path))
    return [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]


def _spacing_pt(length) -> float:
    return length.pt if length is not None else 0.0


def test_build_frontmatter_document_uses_expected_text_order(tmp_path: Path) -> None:
    output_path = tmp_path / "frontmatter.docx"

    build_frontmatter_document(_build_sample_frontmatter(), output_path)

    texts = _normalized_paragraph_texts(output_path)
    expected_order = [
        "First-tier Tribunal\n(Immigration and Asylum Chamber)",
        "Appeal Number: TST-123",
        "THE IMMIGRATION ACTS",
        "Heard at Manchester\nOn 23 January 2026",
        "Before",
        "FIRST-TIER TRIBUNAL JUDGE N SONI",
        "Between",
        "ELLIOTT\n(ANONYMITY ORDER MADE)",
        "Appellant",
        "and",
        "THE SECRETARY OF STATE FOR THE HOME DEPARTMENT",
        "Respondent",
        "Representation:",
    ]

    indices = [texts.index(value) for value in expected_order]
    assert indices == sorted(indices)


def test_frontmatter_uses_requested_underlines_and_label_value_alignment(tmp_path: Path) -> None:
    output_path = tmp_path / "frontmatter.docx"
    build_frontmatter_document(_build_sample_frontmatter(), output_path)
    document = Document(str(output_path))

    acts_heading = next(paragraph for paragraph in document.paragraphs if paragraph.text == "THE IMMIGRATION ACTS")
    assert any(run.underline for run in acts_heading.runs)

    appeal_number_paragraph = next(paragraph for paragraph in document.paragraphs if "Appeal Number:" in paragraph.text)
    assert appeal_number_paragraph.alignment == WD_ALIGN_PARAGRAPH.RIGHT

    representation_heading = next(
        paragraph for paragraph in document.paragraphs if paragraph.text.strip() == "Representation:"
    )
    assert any(run.underline for run in representation_heading.runs)
    assert any(run.bold for run in representation_heading.runs)

    for role_label in ("Appellant", "Respondent"):
        paragraph = next(p for p in document.paragraphs if p.text.strip() == role_label)
        non_space_runs = [run for run in paragraph.runs if run.text.strip()]
        assert non_space_runs
        assert any(run.underline for run in non_space_runs)
        assert all(run.bold is not True for run in non_space_runs)
        assert paragraph.alignment == WD_ALIGN_PARAGRAPH.RIGHT
        assert not any("\t" in run.text for run in paragraph.runs)

    appellant_line = next(p for p in document.paragraphs if p.text.startswith("For the Appellant:"))
    respondent_line = next(p for p in document.paragraphs if p.text.startswith("For the Respondent:"))
    assert len(appellant_line.paragraph_format.tab_stops) == 1
    assert len(respondent_line.paragraph_format.tab_stops) == 1
    assert (
        appellant_line.paragraph_format.tab_stops[0].position == respondent_line.paragraph_format.tab_stops[0].position
    )
    assert "_____" not in appellant_line.text

    judge_line = next(paragraph for paragraph in document.paragraphs if "N SONI" in paragraph.text)
    assert any(run.font.all_caps for run in judge_line.runs if run.text.strip())


def test_frontmatter_omits_anonymity_line_when_disabled(tmp_path: Path) -> None:
    output_path = tmp_path / "frontmatter.docx"

    build_frontmatter_document(_build_sample_frontmatter(include_anonymity_label=False), output_path)

    texts = _normalized_paragraph_texts(output_path)
    assert "ELLIOTT" in texts
    assert "(ANONYMITY ORDER MADE)" not in "\n".join(texts)


def test_frontmatter_uses_blank_lines_and_zero_paragraph_spacing(tmp_path: Path) -> None:
    output_path = tmp_path / "frontmatter.docx"
    build_frontmatter_document(_build_sample_frontmatter(), output_path)
    document = Document(str(output_path))

    blank_lines = [paragraph for paragraph in document.paragraphs if not paragraph.text.strip()]
    assert len(blank_lines) >= 5

    for paragraph in document.paragraphs:
        assert isclose(_spacing_pt(paragraph.paragraph_format.space_before), 0.0, abs_tol=0.01)
        assert isclose(_spacing_pt(paragraph.paragraph_format.space_after), 0.0, abs_tol=0.01)


def test_frontmatter_inserts_requested_blank_paragraph_positions(tmp_path: Path) -> None:
    output_path = tmp_path / "frontmatter.docx"
    build_frontmatter_document(_build_sample_frontmatter(), output_path)
    document = Document(str(output_path))
    paragraph_texts = [paragraph.text.strip() for paragraph in document.paragraphs]

    acts_idx = paragraph_texts.index("THE IMMIGRATION ACTS")
    before_idx = paragraph_texts.index("Before")
    between_idx = paragraph_texts.index("Between")
    and_idx = paragraph_texts.index("and")
    representation_idx = paragraph_texts.index("Representation:")

    assert paragraph_texts[acts_idx - 1] == ""
    assert paragraph_texts[before_idx + 1] == ""
    assert paragraph_texts[between_idx + 1] == ""
    assert paragraph_texts[and_idx + 1] == ""
    assert paragraph_texts[representation_idx + 1] == ""
